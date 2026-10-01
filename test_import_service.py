"""Phase 5G import tests. Every test uses a throwaway SQLite database."""

import csv
import io
import os
import sqlite3
import tempfile
import unittest
from pathlib import Path
from unittest import mock

from database.connection import get_connection, init_database
from database import import_service


class ImportServiceTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.db_path = os.path.join(tempfile.gettempdir(), "pharmacy_import_service_test.db")
        if os.path.exists(cls.db_path):
            os.remove(cls.db_path)
        os.environ["PHARMACY_DB"] = cls.db_path
        init_database()
        conn = get_connection()
        conn.execute("""CREATE TABLE IF NOT EXISTS import_history (
            id INTEGER PRIMARY KEY AUTOINCREMENT, timestamp TEXT NOT NULL, master_type TEXT NOT NULL,
            filename TEXT NOT NULL, total_rows INTEGER NOT NULL, inserted INTEGER NOT NULL,
            updated INTEGER NOT NULL, skipped INTEGER NOT NULL, failed INTEGER NOT NULL)""")
        conn.commit(); conn.close()

    def setUp(self):
        conn = get_connection()
        tables = ["item_ingredients", "items", "doctors", "customers", "suppliers", "drugs", "units", "companies", "account_ledgers", "import_history"]
        for table in tables:
            conn.execute(f"DELETE FROM {table}")
        conn.commit(); conn.close()
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def data(self, headers, rows):
        return {"headers": headers, "rows": [dict(zip(headers, row)) for row in rows]}

    def import_one(self, master, headers, row, mode="skip"):
        preview = import_service.preview_import(master, self.data(headers, [row]))
        return import_service.import_rows(master, preview, duplicate_mode=mode, confirm=True)

    def test_01_csv_basic(self):
        path = Path(self.tmp.name) / "a.csv"; path.write_text("Company Name,Short Name\nAcme,AC\n", encoding="utf-8")
        self.assertEqual(import_service.read_csv(path)["rows"][0]["Company Name"], "Acme")

    def test_02_csv_utf8(self):
        path = Path(self.tmp.name) / "a.csv"; path.write_text("Company Name\nMédic\n", encoding="utf-8")
        self.assertEqual(import_service.read_csv(path)["rows"][0]["Company Name"], "Médic")

    def test_03_csv_bom(self):
        path = Path(self.tmp.name) / "a.csv"; path.write_bytes("\ufeffCompany Name\nAcme\n".encode("utf-8"))
        self.assertEqual(import_service.read_csv(path)["headers"], ["Company Name"])

    def test_04_csv_quoted_fields(self):
        path = Path(self.tmp.name) / "a.csv"; path.write_text('Company Name,Short Name\n"A, B",AB\n', encoding="utf-8")
        self.assertEqual(import_service.read_csv(path)["rows"][0]["Company Name"], "A, B")

    def test_05_csv_empty_rows(self):
        path = Path(self.tmp.name) / "a.csv"; path.write_text("Company Name\n\nAcme\n", encoding="utf-8")
        self.assertEqual(len(import_service.read_csv(path)["rows"]), 1)

    def test_06_csv_malformed_row(self):
        path = Path(self.tmp.name) / "a.csv"; path.write_text("Company Name,Short Name\nAcme\n", encoding="utf-8")
        self.assertIn("__malformed__", import_service.read_csv(path)["rows"][0])

    def test_07_csv_missing_file(self):
        with self.assertRaises(FileNotFoundError): import_service.read_csv("missing.csv")

    def test_08_csv_unsupported_encoding(self):
        path = Path(self.tmp.name) / "a.csv"; path.write_bytes(b"Company Name\n\xff")
        with self.assertRaises(import_service.ImportErrorDetail): import_service.read_csv(path)

    def test_09_xlsx_reading(self):
        path = Path(self.tmp.name) / "a.xlsx"; path.write_bytes(import_service.get_import_template("company", "xlsx"))
        self.assertEqual(import_service.read_excel(path, "Company")["headers"], ["company_name", "short_name"])

    def test_10_xlsx_multiple_sheets(self):
        import openpyxl
        book = openpyxl.Workbook(); book.active.title = "Company"; book.create_sheet("Unit")
        path = Path(self.tmp.name) / "a.xlsx"; book.save(path); book.close()
        self.assertEqual(import_service.read_excel(path)["sheets"], ["Company", "Unit"])

    def test_11_xlsx_selected_sheet(self):
        import openpyxl
        book = openpyxl.Workbook(); book.active.title = "Company"; book.create_sheet("Unit").append(["Unit Name"])
        path = Path(self.tmp.name) / "a.xlsx"; book.save(path); book.close()
        self.assertEqual(import_service.read_excel(path, "Unit")["headers"], ["Unit Name"])

    def test_12_xlsx_missing_sheet(self):
        path = Path(self.tmp.name) / "a.xlsx"; path.write_bytes(import_service.get_import_template("unit", "xlsx"))
        with self.assertRaises(import_service.ImportErrorDetail): import_service.read_excel(path, "Missing")

    def test_13_xlsx_malformed_workbook(self):
        path = Path(self.tmp.name) / "a.xlsx"; path.write_text("not xlsx", encoding="utf-8")
        with self.assertRaises(import_service.ImportErrorDetail): import_service.read_excel(path, "Sheet")

    def test_14_xls_limitation(self):
        with self.assertRaises(import_service.ImportErrorDetail): import_service.inspect_file("file.xls")

    def test_15_inspect_csv(self):
        path = Path(self.tmp.name) / "a.csv"; path.write_text("Unit Name\nBox\n", encoding="utf-8")
        self.assertEqual(import_service.inspect_file(path)["total_rows"] if "total_rows" in import_service.inspect_file(path) else len(import_service.inspect_file(path)["rows"]), 1)

    def test_16_detect_case_insensitive(self):
        self.assertEqual(import_service.detect_columns([" company NAME "], "company"), {" company NAME ": "company_name"})

    def test_17_detect_alias(self):
        self.assertEqual(import_service.detect_columns(["Supplier Name"], "supplier"), {"Supplier Name": "supplier_name"})

    def test_18_manual_mapping(self):
        result = import_service.preview_import("unit", self.data(["Source"], [["Box"]]), {"Source": "unit_name"})
        self.assertEqual(result["valid_rows"], 1)

    def test_19_unmapped_columns(self):
        result = import_service.preview_import("unit", self.data(["Unit Name", "Unused"], [["Box", "x"]]))
        self.assertEqual(result["unmapped_columns"], ["Unused"])

    def test_20_preview_counts(self):
        result = import_service.preview_import("unit", self.data(["Unit Name"], [["A"], ["B"]]))
        self.assertEqual((result["total_rows"], result["valid_rows"]), (2, 2))

    def test_21_preview_read_only(self):
        before = Path(self.db_path).read_bytes()
        import_service.preview_import("unit", self.data(["Unit Name"], [["A"]]))
        self.assertEqual(Path(self.db_path).read_bytes(), before)

    def test_22_required_field(self):
        result = import_service.preview_import("unit", self.data(["Unit Name"], [[""]]))
        self.assertGreater(result["invalid_rows"], 0)

    def test_23_numeric_validation(self):
        result = import_service.preview_import("supplier", self.data(["Supplier Name", "Discount"], [["A", "bad"]]))
        self.assertGreater(result["invalid_rows"], 0)

    def test_24_integer_validation(self):
        result = import_service.preview_import("customer", self.data(["Customer Name", "Credit Period"], [["A", "1.5"]]))
        self.assertGreater(result["invalid_rows"], 0)

    def test_25_boolean_validation(self):
        result = import_service.preview_import("item", self.data(["Item Name", "Scheduled"], [["A", "maybe"]]))
        self.assertGreater(result["invalid_rows"], 0)

    def test_26_blank_rows(self):
        result = import_service.preview_import("unit", {"headers": ["Unit Name"], "rows": []})
        self.assertEqual(result["total_rows"], 0)

    def test_27_duplicate_file_rows(self):
        result = import_service.preview_import("unit", self.data(["Unit Name"], [["A"], ["A"]]))
        self.assertEqual(len(result["duplicate_rows"]), 1)

    def test_28_missing_company_reference(self):
        result = import_service.preview_import("item", self.data(["Item Name", "Company"], [["A", "Missing"]]))
        self.assertGreater(result["invalid_rows"], 0)

    def test_29_missing_unit_reference(self):
        result = import_service.preview_import("item", self.data(["Item Name", "Unit"], [["A", "Missing"]]))
        self.assertGreater(result["invalid_rows"], 0)

    def test_30_missing_drug_reference(self):
        result = import_service.preview_import("item", self.data(["Item Name", "Ingredients"], [["A", "Missing"]]))
        self.assertGreater(result["invalid_rows"], 0)

    def test_31_company_import(self):
        self.assertEqual(self.import_one("company", ["Company Name"], ["Acme"])["inserted"], 1)

    def test_32_unit_import(self):
        self.assertEqual(self.import_one("unit", ["Unit Name"], ["Box"])["inserted"], 1)

    def test_33_drug_import(self):
        self.assertEqual(self.import_one("drug", ["Drug Name"], ["Paracetamol"])["inserted"], 1)

    def test_34_supplier_import(self):
        self.assertEqual(self.import_one("supplier", ["Supplier Name"], ["Acme Supplies"])["inserted"], 1)

    def test_35_customer_import(self):
        self.assertEqual(self.import_one("customer", ["Customer Name"], ["Walk In"])["inserted"], 1)

    def test_36_doctor_import(self):
        self.assertEqual(self.import_one("doctor", ["Doctor Name"], ["Dr Smith"])["inserted"], 1)

    def test_37_item_import(self):
        self.import_one("company", ["Company Name"], ["Acme"]); self.import_one("unit", ["Unit Name"], ["Box"])
        self.assertEqual(self.import_one("item", ["Item Name", "Company", "Unit"], ["Tablet", "Acme", "Box"])["inserted"], 1)

    def test_38_supplier_ledger_created(self):
        self.import_one("supplier", ["Supplier Name"], ["Acme"])
        conn = get_connection(); self.assertEqual(conn.execute("select count(*) from account_ledgers").fetchone()[0], 1); conn.close()

    def test_39_customer_ledger_created(self):
        self.import_one("customer", ["Customer Name"], ["Walk In"])
        conn = get_connection(); self.assertEqual(conn.execute("select count(*) from account_ledgers").fetchone()[0], 1); conn.close()

    def test_40_item_ingredient_import(self):
        self.import_one("unit", ["Unit Name"], ["Box"]); self.import_one("drug", ["Drug Name"], ["Paracetamol"])
        self.import_one("item", ["Item Name", "Unit", "Ingredients"], ["Tablet", "Box", "Paracetamol"])
        conn = get_connection(); self.assertEqual(conn.execute("select count(*) from item_ingredients").fetchone()[0], 1); conn.close()

    def test_41_duplicate_skip(self):
        self.import_one("unit", ["Unit Name"], ["Box"])
        self.assertEqual(self.import_one("unit", ["Unit Name"], ["Box"])["skipped"], 1)

    def test_42_duplicate_update(self):
        self.import_one("doctor", ["Doctor Name", "City"], ["Dr A", "Old"])
        self.assertEqual(self.import_one("doctor", ["Doctor Name", "City"], ["Dr A", "New"], "update")["updated"], 1)

    def test_43_duplicate_cancel(self):
        self.import_one("unit", ["Unit Name"], ["Box"])
        with self.assertRaises(import_service.ImportErrorDetail): self.import_one("unit", ["Unit Name"], ["Box"], "cancel")

    def test_44_confirmation_required(self):
        p = import_service.preview_import("unit", self.data(["Unit Name"], [["Box"]]))
        with self.assertRaises(import_service.ImportErrorDetail): import_service.import_rows("unit", p)

    def test_45_validation_blocks_import(self):
        p = import_service.preview_import("unit", self.data(["Unit Name"], [[""]]))
        with self.assertRaises(import_service.ImportErrorDetail): import_service.import_rows("unit", p, confirm=True)

    def test_46_atomic_rollback(self):
        p = import_service.preview_import("unit", self.data(["Unit Name"], [["A"], ["B"]]))
        with mock.patch.object(import_service, "_write_row", side_effect=[1, RuntimeError("boom")]):
            with self.assertRaises(RuntimeError): import_service.import_rows("unit", p, confirm=True)
        conn = get_connection(); self.assertEqual(conn.execute("select count(*) from units").fetchone()[0], 0); conn.close()

    def test_47_import_commit(self):
        self.import_one("unit", ["Unit Name"], ["Box"])
        conn = get_connection(); self.assertEqual(conn.execute("select count(*) from units").fetchone()[0], 1); conn.close()

    def test_48_persistence_after_reopen(self):
        self.import_one("drug", ["Drug Name"], ["Aspirin"])
        conn = sqlite3.connect(self.db_path); self.assertEqual(conn.execute("select drug_name from drugs").fetchone()[0], "Aspirin"); conn.close()

    def test_49_template_csv(self):
        self.assertIn(b"company_name", import_service.get_import_template("company", "csv"))

    def test_50_template_xlsx(self):
        self.assertTrue(import_service.get_import_template("company", "xlsx").startswith(b"PK"))

    def test_51_template_no_write(self):
        before = Path(self.db_path).read_bytes(); import_service.get_import_template("item", "csv")
        self.assertEqual(Path(self.db_path).read_bytes(), before)

    def test_52_template_all_masters(self):
        for master in import_service.MASTER_TYPES: self.assertTrue(import_service.get_import_template(master, "csv"))

    def test_53_audit_history(self):
        self.import_one("unit", ["Unit Name"], ["Box"])
        self.assertEqual(len(import_service.get_import_history()), 1)

    def test_54_report(self):
        self.assertIn("Inserted: 2", import_service.generate_import_report({"inserted": 2}))

    def test_55_normalize_value(self):
        self.assertEqual(import_service.normalize_value("  A   B \ufeff"), "A B")

    def test_56_invalid_master(self):
        with self.assertRaises(import_service.ImportErrorDetail): import_service.detect_columns([], "invoice")

    def test_57_invalid_extension(self):
        with self.assertRaises(import_service.ImportErrorDetail): import_service.inspect_file("file.txt")

    def test_58_case_insensitive_reference(self):
        self.import_one("unit", ["Unit Name"], ["Box"])
        self.import_one("item", ["Item Name", "Unit"], ["Tablet", " box "])
        conn = get_connection(); self.assertEqual(conn.execute("select count(*) from items").fetchone()[0], 1); conn.close()

    def test_59_numeric_values_persist(self):
        self.import_one("supplier", ["Supplier Name", "Discount"], ["A", "2.5"])
        conn = get_connection(); self.assertEqual(conn.execute("select discount from suppliers").fetchone()[0], 2.5); conn.close()

    def test_60_empty_optional_values_default(self):
        self.import_one("company", ["Company Name"], ["Acme"])
        conn = get_connection(); self.assertEqual(conn.execute("select short_name from companies").fetchone()[0], "Acme"); conn.close()

    def test_61_no_legacy_database_access(self):
        self.assertNotIn("mysql", Path(import_service.__file__).read_text(encoding="utf-8").casefold())

    def test_62_history_contains_counts(self):
        self.import_one("unit", ["Unit Name"], ["Box"])
        history = import_service.get_import_history()[0]
        self.assertEqual((history["inserted"], history["failed"]), (1, 0))

    def test_63_history_empty_is_readable(self):
        self.assertEqual(import_service.get_import_history(), [])

    def test_64_history_multiple_records(self):
        self.import_one("unit", ["Unit Name"], ["Box"])
        self.import_one("drug", ["Drug Name"], ["Aspirin"])
        self.assertEqual(len(import_service.get_import_history()), 2)

    def test_65_history_contents_include_filename_and_counts(self):
        preview = import_service.preview_import("unit", self.data(["Unit Name"], [["Box"]]))
        preview["filename"] = "C:/imports/very-long-unit-master-file.csv"
        import_service.import_rows("unit", preview, confirm=True)
        record = import_service.get_import_history()[0]
        self.assertEqual(record["filename"], preview["filename"])
        self.assertEqual(record["total_rows"], 1)
        self.assertEqual(record["inserted"], 1)

    def test_66_mapping_reports_mapped_and_unmapped_columns(self):
        data = self.data(["Unit Name", "OldFieldX"], [["Box", "legacy"]])
        preview = import_service.preview_import("unit", data)
        self.assertEqual(preview["mapping"], {"Unit Name": "unit_name"})
        self.assertEqual(preview["unmapped_columns"], ["OldFieldX"])

    def test_67_required_target_missing_is_reported(self):
        preview = import_service.preview_import("unit", self.data(["OldFieldX"], [["legacy"]]))
        self.assertTrue(any(issue["field"] == "unit_name" for issue in preview["missing_required_fields"]))

    def test_68_manual_mapping_after_auto_map(self):
        data = self.data(["Unit Name", "Legacy Unit"], [["Box", "Bottle"]])
        auto_mapping = import_service.detect_columns(data["headers"], "unit")
        manual_mapping = dict(auto_mapping)
        manual_mapping["Legacy Unit"] = "unit_name"
        preview = import_service.preview_import("unit", data, manual_mapping)
        self.assertEqual(preview["mapping"]["Legacy Unit"], "unit_name")

    def test_69_preview_mapping_remains_read_only(self):
        before = Path(self.db_path).read_bytes()
        import_service.preview_import("company", self.data(["Company Name", "Old"], [["Acme", "x"]]))
        self.assertEqual(Path(self.db_path).read_bytes(), before)


if __name__ == "__main__":
    unittest.main()
